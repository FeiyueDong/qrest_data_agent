"""XLSX parser: workbook.md + one CSV per sheet (UTF-8, header preserved)."""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path

from qrest_agent.documents.base import (
    DocumentParser,
    ParseError,
    ParseResult,
    fresh_output_dir,
)

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover
    raise ParseError("openpyxl is required for XLSX parsing. Install: pip install openpyxl") from exc

_SHEET_SAFE = re.compile(r'[\\/:*?"<>|\x00-\x1f]+')


def _safe_sheet_name(name: str) -> str:
    safe = _SHEET_SAFE.sub("_", name).strip().replace(" ", "_")
    if not safe:
        safe = "Sheet"
    return safe[:60]


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _sheet_rows(ws) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        text_row = [_cell_text(v) for v in row]
        while text_row and text_row[-1] == "":
            text_row.pop()
        if text_row:
            rows.append(text_row)
    return rows


class XlsxParser(DocumentParser):
    name = "xlsx"
    extensions = (".xlsx", ".xlsm")

    def parse(self, source, output_dir) -> ParseResult:
        src = Path(source)
        if not src.is_file():
            raise ParseError(f"Source file not found: {src}")
        try:
            workbook = load_workbook(str(src), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise ParseError(f"Cannot open XLSX '{src.name}': {exc}") from exc

        out = fresh_output_dir(output_dir)
        md: list[str] = ["# Workbook", "", f"Source: {src.name}", ""]
        sheets_info: list[dict] = []
        csv_names: list[str] = []
        seen_csv: dict[str, int] = {}
        try:
            for sheet_name in workbook.sheetnames:
                ws = workbook[sheet_name]
                rows = _sheet_rows(ws)
                base = _safe_sheet_name(sheet_name)
                seen_csv[base] = seen_csv.get(base, 0) + 1
                if seen_csv[base] > 1:
                    base = f"{base}-{seen_csv[base]}"
                csv_name = base + ".csv"
                csv_names.append(csv_name)
                csv_path = out / csv_name
                with csv_path.open("w", encoding="utf-8", newline="") as fh:
                    writer = csv.writer(fh, lineterminator="\n")
                    writer.writerows(rows)
                width = len(rows[0]) if rows else 0
                headers = rows[0] if rows else []
                sheets_info.append(
                    {
                        "name": sheet_name,
                        "rows": len(rows),
                        "columns": width,
                        "file": csv_name,
                    }
                )
                md.append(f"## Sheet: {sheet_name}")
                md.append("")
                md.append(f"Rows: {len(rows)}")
                md.append(f"Columns: {width}")
                md.append("")
                if headers:
                    md.append("Columns:")
                    md.append("")
                    for header in headers:
                        if header != "":
                            md.append(f"- {header}")
                    md.append("")
                md.append("File:")
                md.append("")
                md.append(csv_name)
                md.append("")
        finally:
            workbook.close()

        workbook_md = out / "workbook.md"
        workbook_md.write_text("\n".join(md).rstrip() + "\n", encoding="utf-8")
        return ParseResult(
            source=src,
            output_files=[workbook_md, *[out / name for name in csv_names]],
            title=None,
            metadata={"sheets": sheets_info},
        )
