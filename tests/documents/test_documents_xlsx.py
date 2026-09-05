"""XLSX parser unit tests."""

from __future__ import annotations

from pathlib import Path

import openpyxl

from qrest_agent.documents import XlsxParser


def _make_xlsx(path: Path) -> None:
    wb = openpyxl.Workbook()
    sensors = wb.active
    sensors.title = "Sensors"
    sensors.append(["SensorID", "Floor", "Direction", "Note"])
    sensors.append(["S001", 1, "X", "首层"])
    sensors.append(["S002", 8, "Y", "中部"])

    channels = wb.create_sheet("Channels")
    channels.append(["ChannelID", "SensorID", "Rate(Hz)"])
    channels.append(["CH-001", "S001", 200])
    wb.save(str(path))


def test_xlsx_workbook_md_and_csv(tmp_path: Path) -> None:
    src = tmp_path / "monitoring.xlsx"
    _make_xlsx(src)
    result = XlsxParser().parse(src, tmp_path / "out")
    assert (tmp_path / "out" / "workbook.md").exists()
    assert (tmp_path / "out" / "Sensors.csv").exists()
    assert (tmp_path / "out" / "Channels.csv").exists()

    sensors_csv = (tmp_path / "out" / "Sensors.csv").read_text(encoding="utf-8")
    lines = sensors_csv.splitlines()
    assert lines[0] == "SensorID,Floor,Direction,Note"
    assert lines[1] == "S001,1,X,首层"

    workbook = (tmp_path / "out" / "workbook.md").read_text(encoding="utf-8")
    assert "## Sheet: Sensors" in workbook
    assert "Rows: 3" in workbook
    assert "Columns: 4" in workbook
    assert "- SensorID" in workbook
    assert "File:" in workbook
    assert "Sensors.csv" in workbook

    sheets = result.metadata["sheets"]
    assert [s["name"] for s in sheets] == ["Sensors", "Channels"]
    assert sheets[0]["rows"] == 3


def test_xlsx_sheet_name_sanitized(tmp_path: Path) -> None:
    src = tmp_path / "odd.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sensors A"
    ws.append(["id"])
    ws.append(["1"])
    wb.save(str(src))
    result = XlsxParser().parse(src, tmp_path / "out")
    assert (tmp_path / "out" / "Sensors_A.csv").exists()
    assert result.metadata["sheets"][0]["file"] == "Sensors_A.csv"
