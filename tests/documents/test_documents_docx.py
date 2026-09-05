"""DOCX parser unit tests (fixtures built in-memory with python-docx)."""

from __future__ import annotations

import json
from pathlib import Path

import docx

from qrest_agent.documents import DocxParser


def _make_docx(path: Path) -> None:
    document = docx.Document()
    document.add_heading("工程概况", 1)
    document.add_paragraph("该建筑位于昆明市，地上14层。")
    document.add_heading("主要参数", 2)
    table = document.add_table(rows=3, cols=2)
    cells = [["参数", "数值"], ["层数", "14"], ["高度", "47.4 m"]]
    for row_index, row in enumerate(cells):
        for col_index, value in enumerate(row):
            table.cell(row_index, col_index).text = value
    document.add_paragraph("监测系统包含28个传感器。")
    document.save(str(path))


def test_docx_headings_tables_order(tmp_path: Path) -> None:
    src = tmp_path / "building.docx"
    _make_docx(src)
    result = DocxParser().parse(src, tmp_path / "out")
    md = (tmp_path / "out" / "document.md").read_text(encoding="utf-8")
    assert result.title == "工程概况"
    assert "# 工程概况" in md
    assert "## 主要参数" in md
    assert "| 层数 | 14 |" in md
    assert "| 高度 | 47.4 m |" in md
    assert "监测系统包含28个传感器。" in md
    # order: heading before table
    assert md.index("工程概况") < md.index("| 参数 | 数值 |")

    source_map = json.loads((tmp_path / "out" / "source_map.json").read_text(encoding="utf-8"))
    kinds = [b["type"] for b in source_map["blocks"]]
    assert kinds == ["heading", "heading", "table"]
    assert result.metadata["headings"] == ["工程概况", "主要参数"]


def test_docx_title_style_is_a_heading(tmp_path: Path) -> None:
    src = tmp_path / "title.docx"
    document = docx.Document()
    document.add_heading("设计说明", 0)
    document.add_paragraph("正文内容。")
    document.save(str(src))
    result = DocxParser().parse(src, tmp_path / "out")
    md = (tmp_path / "out" / "document.md").read_text(encoding="utf-8")
    assert md.startswith("# 设计说明")
    assert result.title == "设计说明"
